# -*- coding: utf-8 -*-
"""
RAG (Retrieval-Augmented Generation) для Zeta.
Читает документы: PDF, Word, Excel, TXT, CSV.
Отвечает на вопросы по содержимому.
Поддерживает: семантический поиск (ChromaDB), поиск по ключевым словам (SQLite fallback).

Особенности:
    - Ленивая инициализация (без падения при импорте)
    - Авто-миграция БД (добавляет недостающие колонки)
    - Автоматическое сохранение в ChromaDB
    - Безопасная разбивка на чанки
    - Прямой вызов Ollama (без рекурсии)
    - Логирование в data/zeta.log
"""

import os
import re
import sys
import hashlib
import sqlite3
import logging
import requests
from pathlib import Path
from typing import List, Dict, Optional, Tuple, Any
from datetime import datetime

sys.path.insert(0, str(Path(__file__).parent.parent))


# ========== ЛОГИРОВАНИЕ ==========

def _log(msg: str) -> None:
    logging.info(f"[RAG] {msg}")


def _log_warn(msg: str) -> None:
    logging.warning(f"[RAG] {msg}")


# ========== КОНСТАНТЫ ==========

DB_PATH = "data/rag.db"
CHROMA_PATH = "data/rag_vectors"
CHROMA_COLLECTION = "documents"

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50
MAX_CONTENT_LENGTH = 3000
MAX_SEARCH_RESULTS = 5

OLLAMA_URL = "http://localhost:11434/api/generate"
OLLAMA_MODEL = "zeta-universal"

# ChromaDB
try:
    import chromadb
    from chromadb.utils import embedding_functions
    HAS_CHROMA = True
except ImportError:
    HAS_CHROMA = False


# ========== ГЛОБАЛЬНОЕ СОСТОЯНИЕ ==========

_db_initialized = False
_collection = None
_collection_ready = False


# ========== ОБРАБОТЧИКИ ДОКУМЕНТОВ ==========

def read_txt(file_path: str) -> str:
    """Читает текстовый файл с автоопределением кодировки."""
    encodings = ["utf-8", "utf-16", "windows-1251", "cp866", "latin-1"]

    for encoding in encodings:
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except Exception:
            continue

    try:
        import chardet
        with open(file_path, "rb") as f:
            raw = f.read()
            result = chardet.detect(raw)
            enc = result.get("encoding", "utf-8")
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
    except ImportError:
        return "⚠️ Установите chardet: pip install chardet"
    except Exception as e:
        return f"⚠️ Не удалось прочитать файл: {e}"


def read_pdf(file_path: str) -> str:
    """Читает PDF-файл."""
    try:
        if os.path.getsize(file_path) < 100:
            return "⚠️ Файл слишком маленький (возможно повреждён)"
    except Exception:
        pass

    try:
        import fitz
        text = ""
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc, 1):
            page_text = page.get_text()
            if page_text.strip():
                text += f"\n--- Страница {page_num} ---\n"
                text += page_text
        doc.close()
        if not text.strip():
            return "⚠️ Не удалось извлечь текст из PDF (возможно, сканированный документ)"
        return text
    except ImportError:
        pass
    except Exception as e:
        _log_warn(f"PyMuPDF ошибка: {e}")

    try:
        import PyPDF2
        text = ""
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            if len(reader.pages) == 0:
                return "⚠️ PDF не содержит страниц"
            for page_num, page in enumerate(reader.pages, 1):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += f"\n--- Страница {page_num} ---\n"
                        text += page_text
                except Exception as e:
                    text += f"\n⚠️ Ошибка на странице {page_num}: {e}\n"
        if not text.strip():
            return "⚠️ Не удалось извлечь текст из PDF"
        return text
    except ImportError:
        return "⚠️ Установите PyMuPDF: pip install PyMuPDF"
    except Exception as e:
        return f"⚠️ Ошибка чтения PDF: {e}"


def read_docx(file_path: str) -> str:
    """Читает Word (.docx)."""
    try:
        import docx
        doc = docx.Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs])
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text for cell in row.cells])
                if row_text.strip():
                    text += "\n" + row_text
        return text if text.strip() else "⚠️ Документ пуст"
    except ImportError:
        return "⚠️ Установите python-docx: pip install python-docx"
    except Exception as e:
        return f"⚠️ Ошибка чтения DOCX: {e}"


def read_excel(file_path: str) -> str:
    """Читает Excel (.xlsx)."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        text = ""
        for sheet in wb.worksheets:
            text += f"\n📋 Лист: {sheet.title}\n"
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(c) if c is not None else "" for c in row])
                if row_text.strip():
                    text += row_text + "\n"
        return text if text.strip() else "⚠️ Excel файл пуст"
    except ImportError:
        return "⚠️ Установите openpyxl: pip install openpyxl"
    except Exception as e:
        return f"⚠️ Ошибка чтения Excel: {e}"


def read_csv(file_path: str) -> str:
    """Читает CSV."""
    try:
        import csv
        text = ""
        with open(file_path, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            for row in reader:
                text += " | ".join(row) + "\n"
        return text if text.strip() else "⚠️ CSV файл пуст"
    except Exception as e:
        return f"⚠️ Ошибка чтения CSV: {e}"


def read_file(file_path: str) -> str:
    """Автоопределение типа файла и чтение."""
    if not os.path.exists(file_path):
        return f"⚠️ Файл не найден: {file_path}"

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".txt":
        return read_txt(file_path)
    elif ext == ".pdf":
        return read_pdf(file_path)
    elif ext == ".docx":
        return read_docx(file_path)
    elif ext in (".xlsx", ".xls"):
        return read_excel(file_path)
    elif ext == ".csv":
        return read_csv(file_path)
    else:
        return f"⚠️ Неподдерживаемый формат: {ext}"


# ========== БАЗА ДАННЫХ (с авто-миграцией) ==========

def _migrate_table(conn: sqlite3.Connection, table: str,
                   expected: Dict[str, str]) -> None:
    """
    Добавляет недостающие колонки в таблицу.
    expected = {"col_name": "TYPE", ...}
    """
    try:
        existing = {
            row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()
        }
    except Exception as e:
        _log_warn(f"PRAGMA {table}: {e}")
        return

    for col, col_type in expected.items():
        if col not in existing:
            try:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} {col_type}")
                _log(f"Миграция: {table}.{col} добавлена")
            except Exception as e:
                _log_warn(f"ALTER {table}.{col}: {e}")


def _ensure_db() -> None:
    """Ленивая инициализация БД с авто-миграцией."""
    global _db_initialized
    if _db_initialized:
        return

    try:
        os.makedirs("data", exist_ok=True)
        with sqlite3.connect(DB_PATH) as conn:
            # --- documents ---
            conn.execute("""
                CREATE TABLE IF NOT EXISTS documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    name TEXT,
                    path TEXT,
                    content TEXT,
                    hash TEXT,
                    size INTEGER,
                    created_at TEXT,
                    updated_at TEXT
                )
            """)

            # --- Авто-миграция ---
            _migrate_table(conn, "documents", {
                "name": "TEXT",
                "path": "TEXT",
                "content": "TEXT",
                "hash": "TEXT",
                "size": "INTEGER",
                "created_at": "TEXT",
                "updated_at": "TEXT",
            })

            # --- chunks ---
            conn.execute("""
                CREATE TABLE IF NOT EXISTS chunks (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    doc_id INTEGER,
                    chunk_text TEXT,
                    chunk_index INTEGER,
                    FOREIGN KEY (doc_id) REFERENCES documents (id)
                )
            """)

            # --- Индексы ---
            conn.execute("CREATE INDEX IF NOT EXISTS idx_chunks_doc_id ON chunks(doc_id)")
            conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_path ON documents(path)")
            conn.execute("CREATE INDEX IF NOT EXISTS idx_documents_name ON documents(name)")

        _db_initialized = True
        _log("БД инициализирована")
    except Exception as e:
        _log_warn(f"Ошибка init_rag_db: {e}")


def init_rag_db() -> None:
    """Публичная функция для инициализации из main.py."""
    _ensure_db()


# ========== CHROMADB ==========

def _get_collection():
    """Ленивая инициализация ChromaDB."""
    global _collection, _collection_ready

    if _collection_ready:
        return _collection

    _collection_ready = True

    if not HAS_CHROMA:
        _collection = None
        return None

    try:
        client = chromadb.PersistentClient(path=CHROMA_PATH)
        _collection = client.get_or_create_collection(
            name=CHROMA_COLLECTION,
            embedding_function=embedding_functions.DefaultEmbeddingFunction(),
        )
        _log("ChromaDB инициализирована")
    except Exception as e:
        _log_warn(f"ChromaDB ошибка: {e}")
        _collection = None

    return _collection


# ========== РАБОТА С ДОКУМЕНТАМИ ==========

def get_file_hash(file_path: str) -> str:
    """Вычисляет хеш файла."""
    try:
        with open(file_path, "rb") as f:
            return hashlib.md5(f.read()).hexdigest()
    except Exception:
        return ""


def split_text(text: str, chunk_size: int = CHUNK_SIZE,
               overlap: int = CHUNK_OVERLAP) -> List[str]:
    """Разбивает текст на чанки."""
    if not text:
        return []

    if chunk_size <= overlap:
        overlap = max(0, chunk_size // 4)

    text = re.sub(r"\s+", " ", text).strip()
    words = text.split()

    if len(words) <= chunk_size:
        return [text]

    step = chunk_size - overlap
    if step <= 0:
        step = chunk_size

    chunks = []
    for i in range(0, len(words), step):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)

    return chunks


def save_document(file_path: str, content: str) -> bool:
    """Сохраняет документ в SQLite + ChromaDB."""
    _ensure_db()

    if not content or len(content) < 10:
        return False

    file_hash = get_file_hash(file_path)
    file_name = os.path.basename(file_path)

    try:
        file_size = os.path.getsize(file_path)
    except Exception:
        file_size = len(content)

    with sqlite3.connect(DB_PATH) as conn:
        existing = conn.execute(
            "SELECT id, hash FROM documents WHERE path = ?", (file_path,)
        ).fetchone()

        if existing:
            doc_id, old_hash = existing
            if old_hash == file_hash:
                _log(f"Документ не изменился: {file_name}")
                return True

            conn.execute("""
                UPDATE documents SET content = ?, hash = ?, size = ?, updated_at = ?
                WHERE id = ?
            """, (content, file_hash, file_size, datetime.now().isoformat(), doc_id))
            conn.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
        else:
            cursor = conn.execute("""
                INSERT INTO documents (name, path, content, hash, size, created_at, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, (file_name, file_path, content, file_hash, file_size,
                  datetime.now().isoformat(), datetime.now().isoformat()))
            doc_id = cursor.lastrowid

        chunks = split_text(content)
        for i, chunk in enumerate(chunks):
            conn.execute("""
                INSERT INTO chunks (doc_id, chunk_text, chunk_index)
                VALUES (?, ?, ?)
            """, (doc_id, chunk, i))

    _save_to_chroma(doc_id, chunks)

    _log(f"Сохранён: {file_name} ({len(chunks)} чанков)")
    return True


def _save_to_chroma(doc_id: int, chunks: List[str]) -> None:
    """Сохраняет чанки в ChromaDB."""
    collection = _get_collection()
    if not collection:
        return

    try:
        try:
            collection.delete(where={"doc_id": doc_id})
        except Exception:
            pass

        ids = [f"doc_{doc_id}_chunk_{i}" for i in range(len(chunks))]
        collection.add(
            ids=ids,
            documents=chunks,
            metadatas=[{"doc_id": doc_id} for _ in chunks],
        )
        _log(f"ChromaDB: +{len(chunks)} чанков для doc_id={doc_id}")
    except Exception as e:
        _log_warn(f"Ошибка сохранения в ChromaDB: {e}")


# ========== ПОИСК ==========

def _search_chroma(query: str, limit: int = MAX_SEARCH_RESULTS) -> List[Dict]:
    """Семантический поиск в ChromaDB."""
    collection = _get_collection()
    if not collection:
        return []

    try:
        results = collection.query(query_texts=[query], n_results=limit)

        documents = results.get("documents") or []
        if not documents or not documents[0]:
            return []

        docs = []
        metadatas = results.get("metadatas") or [[]]
        distances = results.get("distances") or [[]]

        for i, doc_text in enumerate(documents[0]):
            metadata = metadatas[0][i] if metadatas and metadatas[0] else {}
            score = 0.0
            if distances and distances[0] and i < len(distances[0]):
                score = 1.0 - distances[0][i]
            docs.append({
                "content": doc_text,
                "doc_id": metadata.get("doc_id", 0),
                "score": score,
                "preview": doc_text[:300],
            })
        return docs
    except Exception as e:
        _log_warn(f"Ошибка поиска в ChromaDB: {e}")
        return []


def _search_keywords(query: str, limit: int = MAX_SEARCH_RESULTS) -> List[Dict]:
    """Поиск по ключевым словам в chunks."""
    _ensure_db()

    if not query or not query.strip():
        return []

    keywords = [k for k in query.lower().split() if len(k) > 2]
    if not keywords:
        return []

    with sqlite3.connect(DB_PATH) as conn:
        chunks = conn.execute("""
            SELECT c.id, c.doc_id, c.chunk_text, d.name
            FROM chunks c
            JOIN documents d ON d.id = c.doc_id
        """).fetchall()

    results = []
    for chunk_id, doc_id, chunk_text, doc_name in chunks:
        text_lower = chunk_text.lower()
        score = 0
        matched = []
        for kw in keywords:
            count = text_lower.count(kw)
            if count > 0:
                score += count
                matched.append(kw)

        if score > 0:
            results.append({
                "doc_id": doc_id,
                "name": doc_name,
                "score": score,
                "matched_keywords": matched,
                "preview": _get_relevant_preview(chunk_text, keywords),
                "content": chunk_text[:MAX_CONTENT_LENGTH],
            })

    results.sort(key=lambda x: x["score"], reverse=True)
    return results[:limit]


def _get_relevant_preview(text: str, keywords: List[str], max_len: int = 300) -> str:
    """Возвращает релевантный фрагмент."""
    if not text:
        return ""

    text_lower = text.lower()
    for kw in keywords:
        pos = text_lower.find(kw)
        if pos != -1:
            start = max(0, pos - 50)
            end = min(len(text), pos + len(kw) + 50)
            preview = text[start:end]
            if start > 0:
                preview = "..." + preview
            if end < len(text):
                preview = preview + "..."
            if len(preview) > max_len:
                preview = preview[:max_len] + "..."
            return preview

    return text[:max_len] + "..." if len(text) > max_len else text


def hybrid_search(query: str, limit: int = MAX_SEARCH_RESULTS) -> List[Dict]:
    """Комбинирует векторный и ключевой поиск."""
    results = []

    if HAS_CHROMA:
        for r in _search_chroma(query, limit):
            results.append({
                "doc_id": r.get("doc_id", 0),
                "name": f"Chunk {r.get('doc_id', 0)}",
                "score": r.get("score", 0),
                "matched_keywords": [],
                "preview": r.get("preview", ""),
                "content": r.get("content", ""),
                "source": "vector",
            })

    for r in _search_keywords(query, limit):
        results.append({
            "doc_id": r["doc_id"],
            "name": r["name"],
            "score": r["score"],
            "matched_keywords": r["matched_keywords"],
            "preview": r["preview"],
            "content": r["content"],
            "source": "keyword",
        })

    results.sort(key=lambda x: x["score"], reverse=True)

    seen = set()
    unique = []
    for r in results:
        if r["doc_id"] not in seen:
            unique.append(r)
            seen.add(r["doc_id"])

    return unique[:limit]


# ========== RAG-ОТВЕТЫ ==========

def _ask_ollama(prompt: str, timeout: int = 120) -> str:
    """Прямой вызов Ollama."""
    try:
        response = requests.post(
            OLLAMA_URL,
            json={
                "model": OLLAMA_MODEL,
                "prompt": prompt,
                "stream": False,
                "temperature": 0.3,
                "num_predict": 2048,
            },
            timeout=timeout,
        )
        if response.status_code != 200:
            return f"⚠️ Ollama вернул {response.status_code}"

        data = response.json()
        return data.get("response", "").strip() or "(пустой ответ)"
    except requests.exceptions.ConnectionError:
        return "⚠️ Ollama не запущена. Запустите: ollama serve"
    except requests.exceptions.Timeout:
        return "⚠️ Ollama думает слишком долго"
    except Exception as e:
        return f"⚠️ Ошибка: {e}"


def ask_with_rag(query: str) -> str:
    """Отвечает на вопрос, используя документы."""
    if not query or not query.strip():
        return "⚠️ Укажите вопрос для поиска."

    results = hybrid_search(query, limit=MAX_SEARCH_RESULTS)

    if not results:
        return (
            "📭 В базе нет подходящих документов. "
            "Загрузите файл: `загрузи документ D:\\путь\\файл.pdf`"
        )

    context = "📄 Найденные документы:\n\n"
    for i, res in enumerate(results, 1):
        if res.get("source") == "vector":
            context += f"[{i}] (векторный поиск)\n"
        else:
            context += f"[{i}] {res.get('name', 'Документ')} (релевантность: {res['score']})\n"
        context += f"{res['preview']}\n\n"

    prompt = (
        "Ты — Zeta. Тебя создал Самир (Samriddin).\n"
        "Ты отвечаешь на вопрос пользователя ТОЛЬКО на русском языке.\n"
        "Используй информацию из документов ниже. "
        "Если информации нет — честно скажи об этом.\n\n"
        f"=== ДОКУМЕНТЫ ===\n{context}\n"
        f"=== ВОПРОС ===\n{query}\n\n"
        "Ответ:"
    )

    return _ask_ollama(prompt)


# ========== УПРАВЛЕНИЕ ДОКУМЕНТАМИ ==========

def get_all_documents() -> List[Dict]:
    """Список всех документов."""
    _ensure_db()
    with sqlite3.connect(DB_PATH) as conn:
        rows = conn.execute(
            "SELECT id, name, path, size, created_at FROM documents ORDER BY created_at DESC"
        ).fetchall()
    return [
        {"id": r[0], "name": r[1], "path": r[2], "size": r[3], "created_at": r[4]}
        for r in rows
    ]


def get_document(doc_id: int) -> Optional[Dict]:
    """Документ по ID."""
    _ensure_db()
    with sqlite3.connect(DB_PATH) as conn:
        row = conn.execute(
            "SELECT id, name, path, content, size, created_at FROM documents WHERE id = ?",
            (doc_id,),
        ).fetchone()
    if row:
        return {
            "id": row[0], "name": row[1], "path": row[2],
            "content": row[3], "size": row[4], "created_at": row[5],
        }
    return None


def delete_document(doc_id: int) -> bool:
    """Удаляет документ."""
    _ensure_db()
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))

    collection = _get_collection()
    if collection:
        try:
            collection.delete(where={"doc_id": doc_id})
        except Exception as e:
            _log_warn(f"ChromaDB delete: {e}")

    _log(f"Удалён документ id={doc_id}")
    return True


def clear_all_documents() -> None:
    """Удаляет все документы."""
    _ensure_db()
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("DELETE FROM chunks")
        conn.execute("DELETE FROM documents")

    global _collection
    if HAS_CHROMA:
        try:
            client = chromadb.PersistentClient(path=CHROMA_PATH)
            try:
                client.delete_collection(CHROMA_COLLECTION)
            except Exception:
                pass
            _collection = client.get_or_create_collection(
                name=CHROMA_COLLECTION,
                embedding_function=embedding_functions.DefaultEmbeddingFunction(),
            )
            _log("ChromaDB очищена")
        except Exception as e:
            _log_warn(f"ChromaDB clear: {e}")

    _log("Все документы удалены")


def get_document_stats() -> Dict[str, Any]:
    """Статистика документов."""
    _ensure_db()
    with sqlite3.connect(DB_PATH) as conn:
        total = conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        total_chunks = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        total_size = conn.execute("SELECT SUM(size) FROM documents").fetchone()[0] or 0
        total_chars = conn.execute("SELECT SUM(LENGTH(content)) FROM documents").fetchone()[0] or 0

    return {
        "total_documents": total,
        "total_chunks": total_chunks,
        "total_size": total_size,
        "total_characters": total_chars,
    }


# ========== ЭКСПОРТ ==========

__all__ = [
    "read_file",
    "read_txt", "read_pdf", "read_docx", "read_excel", "read_csv",
    "init_rag_db",
    "save_document", "delete_document", "clear_all_documents",
    "get_all_documents", "get_document", "get_document_stats",
    "hybrid_search", "ask_with_rag",
    "split_text",
]


# ========== ТЕСТ ==========

if __name__ == "__main__":
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
    )

    print("🧠 Тест RAG:\n")
    print("=" * 60)

    init_rag_db()

    # --- Тест 1 ---
    print("\n📝 Тест 1: Чтение и сохранение")
    test_file = "test_rag.txt"
    with open(test_file, "w", encoding="utf-8") as f:
        f.write(
            "Zeta — это ИИ-помощник, созданный в Узбекистане. "
            "Создатель — Samriddin (Самриддин). "
            "Zeta умеет отвечать на вопросы, писать код и анализировать изображения. "
            "Проект написан на Python + PyQt6 + Ollama."
        )

    content = read_file(test_file)
    print(f"📄 Прочитано: {len(content)} символов")

    save_document(test_file, content)
    print("✅ Документ сохранён")

    # --- Тест 2 ---
    print("\n📝 Тест 2: Статистика")
    stats = get_document_stats()
    print(f"Документов: {stats['total_documents']}")
    print(f"Чанков: {stats['total_chunks']}")

    # --- Тест 3 ---
    print("\n📝 Тест 3: Гибридный поиск")
    results = hybrid_search("кто создал Zeta")
    for r in results:
        print(f"  • [{r.get('source')}] {r.get('name', 'Chunk')} (score: {r.get('score')})")
        print(f"    {r['preview'][:100]}...")

    # --- Тест 4 ---
    print("\n📝 Тест 4: RAG-ответ")
    answer = ask_with_rag("Кто создал Zeta?")
    print(f"Ответ: {answer[:300]}")

    # --- Очистка ---
    print("\n🧹 Очистка")
    try:
        os.remove(test_file)
    except Exception:
        pass
    clear_all_documents()
    print("✅ Тесты завершены!")